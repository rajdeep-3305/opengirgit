"""Resume text extraction (PDF / DOCX / TXT) — all local, no network."""

from __future__ import annotations

import hashlib
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".rtf"}


class ResumeExtractionError(Exception):
    pass


def extract_text(path: str | Path) -> str:
    path = Path(path)
    if not path.exists():
        raise ResumeExtractionError(f"file not found: {path}")
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in (".txt", ".md", ".rtf"):
        return path.read_text(encoding="utf-8", errors="replace")

    # Unknown extension: try PDF magic, then plain text.
    head = path.read_bytes()[:5]
    if head.startswith(b"%PDF"):
        return _extract_pdf(path)
    try:
        return path.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as exc:
        raise ResumeExtractionError(f"unsupported file type: {suffix or 'unknown'}") from exc


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency hint
        raise ResumeExtractionError("pypdf is not installed (pip install pypdf)") from exc

    try:
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ResumeExtractionError(f"could not read PDF: {exc}") from exc
    text = "\n".join(pages).strip()
    if not text:
        raise ResumeExtractionError("no text found in the PDF (is it a scan?)")
    return text


def _extract_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency hint
        raise ResumeExtractionError("python-docx is not installed (pip install python-docx)") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ResumeExtractionError(f"could not read DOCX: {exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ResumeExtractionError("no text found in the DOCX")
    return text


def text_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", "replace")).hexdigest()[:16]
